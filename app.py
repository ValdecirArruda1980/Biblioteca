import json
import os
import sqlite3
from datetime import datetime
from docx import Document
from flask import Flask, jsonify, render_template, request
import pandas as pd
from pypdf import PdfReader

app = Flask(__name__)
DB_NAME = 'database.db'


def get_db_connection():
  conn = sqlite3.connect(DB_NAME)
  conn.row_factory = sqlite3.Row
  return conn


def init_db():
  conn = get_db_connection()
  cursor = conn.cursor()
  cursor.execute("""
        CREATE TABLE IF NOT EXISTS leitores (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT NOT NULL,
            ra TEXT UNIQUE NOT NULL,
            curso TEXT NOT NULL,
            categoria TEXT DEFAULT 'Aluno'
        );
    """)

  cursor.execute("""
        CREATE TABLE IF NOT EXISTS livros (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            titulo TEXT NOT NULL,
            autor TEXT NOT NULL,
            isbn TEXT,
            editora TEXT,
            ano_publicacao TEXT,
            quantidade INTEGER DEFAULT 1,
            disponiveis INTEGER DEFAULT 1,
            data_cadastro TEXT
        );
    """)

  cursor.execute('PRAGMA table_info(livros)')
  columns = [col['name'] for col in cursor.fetchall()]
  if 'editora' not in columns:
    cursor.execute('ALTER TABLE livros ADD COLUMN editora TEXT')
  if 'ano_publicacao' not in columns:
    cursor.execute('ALTER TABLE livros ADD COLUMN ano_publicacao TEXT')
  if 'data_cadastro' not in columns:
    cursor.execute('ALTER TABLE livros ADD COLUMN data_cadastro TEXT')

  cursor.execute("""
        CREATE TABLE IF NOT EXISTS emprestimos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            leitor_id INTEGER NOT NULL,
            livro_id INTEGER NOT NULL,
            data_emprestimo TEXT NOT NULL,
            data_previsao TEXT NOT NULL,
            data_devolucao TEXT,
            status TEXT DEFAULT 'Ativo',
            FOREIGN KEY (leitor_id) REFERENCES leitores (id),
            FOREIGN KEY (livro_id) REFERENCES livros (id)
        );
    """)
  conn.commit()
  conn.close()


# --- EXTRAÇÃO DE LIVROS DE ARQUIVO ---
def extrair_livros_de_arquivo(file, filename):
  livros_extraidos = []
  ext = filename.split('.')[-1].lower()

  if ext in ['xlsx', 'xls', 'csv']:
    df = pd.read_excel(file) if ext in ['xlsx', 'xls'] else pd.read_csv(file)
    df.columns = [str(c).strip().lower() for c in df.columns]

    for _, row in df.iterrows():
      titulo = str(row.get('titulo', '')).strip()
      autor = str(row.get('autor', 'Autor Desconhecido')).strip()
      isbn = (
          str(row.get('isbn', '')).strip() if pd.notna(row.get('isbn')) else ''
      )
      editora = (
          str(row.get('editora', '')).strip()
          if pd.notna(row.get('editora'))
          else ''
      )
      ano = (
          str(row.get('ano', row.get('ano_publicacao', ''))).strip()
          if pd.notna(row.get('ano', row.get('ano_publicacao')))
          else ''
      )

      try:
        qtd = int(row.get('quantidade', 1))
      except:
        qtd = 1

      if titulo and titulo.lower() != 'nan':
        livros_extraidos.append({
            'titulo': titulo,
            'autor': autor,
            'isbn': isbn,
            'editora': editora,
            'ano_publicacao': ano,
            'quantidade': qtd,
        })

  elif ext == 'docx':
    doc = Document(file)
    for table in doc.tables:
      for row in table.rows[1:]:
        cols = [cell.text.strip() for cell in row.cells]
        if len(cols) >= 1 and cols[0]:
          livros_extraidos.append({
              'titulo': cols[0],
              'autor': cols[1] if len(cols) > 1 else 'Autor Desconhecido',
              'isbn': cols[2] if len(cols) > 2 else '',
              'editora': cols[3] if len(cols) > 3 else '',
              'ano_publicacao': cols[4] if len(cols) > 4 else '',
              'quantidade': int(cols[5])
              if len(cols) > 5 and cols[5].isdigit()
              else 1,
          })

  elif ext == 'pdf':
    reader = PdfReader(file)
    for page in reader.pages:
      texto = page.extract_text()
      if texto:
        for linha in texto.split('\n'):
          linha = linha.strip()
          if '-' in linha and len(linha) > 3:
            partes = linha.split('-')
            livros_extraidos.append({
                'titulo': partes[0].strip(),
                'autor': partes[1].strip()
                if len(partes) > 1
                else 'Desconhecido',
                'isbn': '',
                'editora': '',
                'ano_publicacao': '',
                'quantidade': 1,
            })

  elif ext == 'json':
    data = json.load(file)
    if isinstance(data, list):
      for item in data:
        if 'titulo' in item:
          livros_extraidos.append({
              'titulo': item.get('titulo'),
              'autor': item.get('autor', 'Desconhecido'),
              'isbn': item.get('isbn', ''),
              'editora': item.get('editora', ''),
              'ano_publicacao': item.get('ano_publicacao', ''),
              'quantidade': int(item.get('quantidade', 1)),
          })

  elif ext == 'txt':
    linhas = file.read().decode('utf-8').splitlines()
    for linha in linhas:
      linha = linha.strip()
      if '-' in linha:
        partes = linha.split('-')
        livros_extraidos.append({
            'titulo': partes[0].strip(),
            'autor': partes[1].strip() if len(partes) > 1 else 'Desconhecido',
            'isbn': '',
            'editora': '',
            'ano_publicacao': '',
            'quantidade': 1,
        })

  return livros_extraidos


# --- EXTRAÇÃO DE ALUNOS/LEITORES DE ARQUIVO ---
def extrair_alunos_de_arquivo(file, filename):
  alunos_extraidos = []
  ext = filename.split('.')[-1].lower()

  if ext in ['xlsx', 'xls', 'csv']:
    df = pd.read_excel(file) if ext in ['xlsx', 'xls'] else pd.read_csv(file)
    df.columns = [str(c).strip().lower() for c in df.columns]

    for _, row in df.iterrows():
      nome = str(row.get('nome', '')).strip()
      ra = str(row.get('ra', '')).strip()
      curso = (
          str(row.get('curso', row.get('turma', 'Geral'))).strip()
          if pd.notna(row.get('curso', row.get('turma')))
          else 'Geral'
      )
      categoria = str(row.get('categoria', 'Aluno')).strip()

      if nome and ra and nome.lower() != 'nan' and ra.lower() != 'nan':
        alunos_extraidos.append({
            'nome': nome,
            'ra': ra,
            'curso': curso,
            'categoria': categoria if categoria else 'Aluno',
        })

  elif ext == 'docx':
    doc = Document(file)
    for table in doc.tables:
      for row in table.rows[1:]:
        cols = [cell.text.strip() for cell in row.cells]
        if len(cols) >= 2 and cols[0] and cols[1]:
          alunos_extraidos.append({
              'nome': cols[0],
              'ra': cols[1],
              'curso': cols[2] if len(cols) > 2 else 'Geral',
              'categoria': cols[3] if len(cols) > 3 else 'Aluno',
          })

  elif ext == 'pdf':
    reader = PdfReader(file)
    for page in reader.pages:
      texto = page.extract_text()
      if texto:
        for linha in texto.split('\n'):
          linha = linha.strip()
          if '-' in linha and len(linha) > 3:
            partes = linha.split('-')
            if len(partes) >= 2:
              alunos_extraidos.append({
                  'nome': partes[0].strip(),
                  'ra': partes[1].strip(),
                  'curso': partes[2].strip() if len(partes) > 2 else 'Geral',
                  'categoria': 'Aluno',
              })

  elif ext == 'json':
    data = json.load(file)
    if isinstance(data, list):
      for item in data:
        if 'nome' in item and 'ra' in item:
          alunos_extraidos.append({
              'nome': item.get('nome'),
              'ra': str(item.get('ra')),
              'curso': item.get('curso', 'Geral'),
              'categoria': item.get('categoria', 'Aluno'),
          })

  elif ext == 'txt':
    linhas = file.read().decode('utf-8').splitlines()
    for linha in linhas:
      linha = linha.strip()
      if '-' in linha:
        partes = linha.split('-')
        if len(partes) >= 2:
          alunos_extraidos.append({
              'nome': partes[0].strip(),
              'ra': partes[1].strip(),
              'curso': partes[2].strip() if len(partes) > 2 else 'Geral',
              'categoria': 'Aluno',
          })

  return alunos_extraidos


# --- ROTAS DA APLICAÇÃO ---


@app.route('/')
def index():
  return render_template('index.html')


@app.route('/api/leitores', methods=['GET', 'POST'])
def api_leitores():
  conn = get_db_connection()
  if request.method == 'POST':
    data = request.json
    try:
      conn.execute(
          'INSERT INTO leitores (nome, ra, curso, categoria) VALUES (?, ?, ?,'
          ' ?)',
          (
              data['nome'].strip(),
              data['ra'].strip(),
              data['curso'].strip(),
              data.get('categoria', 'Aluno'),
          ),
      )
      conn.commit()
      return jsonify({'message': 'Leitor cadastrado com sucesso!'}), 201
    except sqlite3.IntegrityError:
      return jsonify({'error': 'RA já cadastrado.'}), 400
    finally:
      conn.close()

  leitores = conn.execute('SELECT * FROM leitores ORDER BY id DESC').fetchall()
  conn.close()
  return jsonify([dict(row) for row in leitores])


@app.route('/api/leitores/importar', methods=['POST'])
def api_importar_leitores():
  if 'file' not in request.files:
    return jsonify({'error': 'Nenhum arquivo enviado.'}), 400

  file = request.files['file']
  if file.filename == '':
    return jsonify({'error': 'Arquivo não selecionado.'}), 400

  try:
    alunos = extrair_alunos_de_arquivo(file, file.filename)

    if not alunos:
      return (
          jsonify({
              'error': (
                  'Não foi possível extrair alunos do arquivo. Verifique se o'
                  ' formato tem as colunas Nome, RA e Curso.'
              )
          }),
          400,
      )

    conn = get_db_connection()
    inseridos = 0
    duplicados = 0

    for a in alunos:
      try:
        conn.execute(
            'INSERT INTO leitores (nome, ra, curso, categoria) VALUES (?, ?,'
            ' ?, ?)',
            (a['nome'], a['ra'], a['curso'], a['categoria']),
        )
        inseridos += 1
      except sqlite3.IntegrityError:
        duplicados += 1

    conn.commit()
    conn.close()

    msg = f'Sucesso! {inseridos} aluno(s) cadastrado(s) com sucesso.'
    if duplicados > 0:
      msg += f' ({duplicados} RA(s) ignorado(s) por já existirem no banco).'

    return jsonify({'message': msg}), 201

  except Exception as e:
    return jsonify({'error': f'Erro ao processar arquivo: {str(e)}'}), 500


@app.route('/api/livros', methods=['GET', 'POST'])
def api_livros():
  conn = get_db_connection()
  if request.method == 'POST':
    data = request.json
    qtd = int(data.get('quantidade', 1))
    hoje = datetime.now().strftime('%d/%m/%Y')

    conn.execute(
        """
            INSERT INTO livros (titulo, autor, isbn, editora, ano_publicacao, quantidade, disponiveis, data_cadastro)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            data['titulo'].strip(),
            data['autor'].strip(),
            data.get('isbn', '').strip(),
            data.get('editora', '').strip(),
            data.get('ano_publicacao', '').strip(),
            qtd,
            qtd,
            hoje,
        ),
    )
    conn.commit()
    conn.close()
    return jsonify({'message': 'Livro cadastrado com sucesso!'}), 201

  livros = conn.execute('SELECT * FROM livros ORDER BY id DESC').fetchall()
  conn.close()
  return jsonify([dict(row) for row in livros])


@app.route('/api/livros/importar', methods=['POST'])
def api_importar_livros():
  if 'file' not in request.files:
    return jsonify({'error': 'Nenhum arquivo enviado.'}), 400

  file = request.files['file']
  if file.filename == '':
    return jsonify({'error': 'Arquivo não selecionado.'}), 400

  try:
    livros = extrair_livros_de_arquivo(file, file.filename)

    if not livros:
      return (
          jsonify({
              'error': (
                  'Não foi possível extrair livros do arquivo. Verifique se o'
                  ' formato está correto.'
              )
          }),
          400,
      )

    conn = get_db_connection()
    hoje = datetime.now().strftime('%d/%m/%Y')
    inseridos = 0

    for l in livros:
      conn.execute(
          """
                INSERT INTO livros (titulo, autor, isbn, editora, ano_publicacao, quantidade, disponiveis, data_cadastro)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
          (
              l['titulo'],
              l['autor'],
              l.get('isbn', ''),
              l.get('editora', ''),
              l.get('ano_publicacao', ''),
              l['quantidade'],
              l['quantidade'],
              hoje,
          ),
      )
      inseridos += 1

    conn.commit()
    conn.close()

    return (
        jsonify({
            'message': (
                f'Sucesso! {inseridos} livro(s) importado(s) com sucesso a'
                ' partir do arquivo.'
            )
        }),
        201,
    )

  except Exception as e:
    return jsonify({'error': f'Erro ao processar o arquivo: {str(e)}'}), 500


@app.route('/api/emprestimos', methods=['GET', 'POST'])
def api_emprestimos():
  conn = get_db_connection()
  if request.method == 'POST':
    data = request.json
    leitor_id = data['leitor_id']
    livro_id = data['livro_id']

    livro = conn.execute(
        'SELECT disponiveis FROM livros WHERE id = ?', (livro_id,)
    ).fetchone()
    if not livro or livro['disponiveis'] < 1:
      conn.close()
      return jsonify({'error': 'Exemplar indisponível para empréstimo.'}), 400

    hoje = datetime.now().strftime('%Y-%m-%d')
    conn.execute(
        'INSERT INTO emprestimos (leitor_id, livro_id, data_emprestimo,'
        " data_previsao, status) VALUES (?, ?, ?, ?, 'Ativo')",
        (leitor_id, livro_id, hoje, data['data_previsao']),
    )

    conn.execute(
        'UPDATE livros SET disponiveis = disponiveis - 1 WHERE id = ?',
        (livro_id,),
    )
    conn.commit()
    conn.close()
    return jsonify({'message': 'Empréstimo registrado com sucesso!'}), 201

  query = """
        SELECT e.id, l.nome as leitor, b.titulo as livro, e.data_emprestimo, e.data_previsao, e.data_devolucao, e.status
        FROM emprestimos e
        JOIN leitores l ON e.leitor_id = l.id
        JOIN livros b ON e.livro_id = b.id
        ORDER BY e.id DESC
    """
  emprestimos = conn.execute(query).fetchall()
  conn.close()
  return jsonify([dict(row) for row in emprestimos])


@app.route('/api/emprestimos/baixa/<int:emprestimo_id>', methods=['POST'])
def api_baixa_emprestimo(emprestimo_id):
  conn = get_db_connection()
  emprestimo = conn.execute(
      'SELECT livro_id, status FROM emprestimos WHERE id = ?', (emprestimo_id,)
  ).fetchone()

  if not emprestimo:
    conn.close()
    return jsonify({'error': 'Empréstimo não encontrado.'}), 404

  if emprestimo['status'] == 'Devolvido':
    conn.close()
    return jsonify({'message': 'Este empréstimo já foi devolvido.'}), 200

  hoje = datetime.now().strftime('%Y-%m-%d')
  conn.execute(
      "UPDATE emprestimos SET status = 'Devolvido', data_devolucao = ? WHERE"
      ' id = ?',
      (hoje, emprestimo_id),
  )
  conn.execute(
      'UPDATE livros SET disponiveis = disponiveis + 1 WHERE id = ?',
      (emprestimo['livro_id'],),
  )

  conn.commit()
  conn.close()
  return jsonify({'message': 'Devolução realizada com sucesso!'})


if __name__ == '__main__':
  init_db()
  app.run(debug=True, host='0.0.0.0', port=5000)
